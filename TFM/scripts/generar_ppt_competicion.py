"""
Genera la presentación .pptx para la competición de becas del TFM.
Motor de recomendación turística con IA y redistribución de demanda (TUI).

Uso:
    python scripts/generar_ppt_competicion.py

Salida:
    docs/Presentacion_Competicion_Becas.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# --- Paleta TUI ---------------------------------------------------------------
TUI_BLUE = RGBColor(0x00, 0x2D, 0x72)      # azul corporativo
TUI_ACCENT = RGBColor(0xE6, 0x00, 0x78)    # magenta acento
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
GREY = RGBColor(0x55, 0x5A, 0x66)
LIGHT = RGBColor(0xF2, 0xF4, 0xF8)

FONT = "Calibri"

ASSETS = Path(__file__).resolve().parent.parent / "docs" / "assets_ppt"

prs = Presentation()
prs.slide_width = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

BLANK = prs.slide_layouts[6]


def add_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_band(slide, top, height, color):
    shape = slide.shapes.add_shape(1, 0, top, SW, height)  # rectangle
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def set_run(run, text, size, color, bold=False, italic=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = FONT


def title_slide(title, subtitle, foot):
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide, TUI_BLUE)
    add_band(slide, Emu(int(SH * 0.62)), Emu(int(SH * 0.012)), TUI_ACCENT)

    _, tf = textbox(slide, Inches(0.9), Inches(2.1), Inches(11.5), Inches(2.4))
    p = tf.paragraphs[0]
    set_run(p.add_run(), title, 40, WHITE, bold=True)
    p.alignment = PP_ALIGN.LEFT

    p2 = tf.add_paragraph()
    set_run(p2.add_run(), subtitle, 22, RGBColor(0xC9, 0xD6, 0xF0))
    p2.space_before = Pt(14)

    _, tf3 = textbox(slide, Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.6))
    p3 = tf3.paragraphs[0]
    set_run(p3.add_run(), foot, 14, RGBColor(0x9F, 0xB2, 0xD8))
    return slide


def content_slide(number, kicker, title, bullets, footer_note=None):
    """bullets: list of (text, level, bold) tuples."""
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide, WHITE)

    # cabecera
    add_band(slide, 0, Inches(1.35), TUI_BLUE)
    # numero
    _, tfn = textbox(slide, Inches(11.9), Inches(0.28), Inches(1.2), Inches(0.8))
    pn = tfn.paragraphs[0]
    pn.alignment = PP_ALIGN.RIGHT
    set_run(pn.add_run(), str(number), 30, RGBColor(0x4A, 0x63, 0x9E), bold=True)

    # kicker
    _, tfk = textbox(slide, Inches(0.6), Inches(0.22), Inches(11), Inches(0.4))
    pk = tfk.paragraphs[0]
    set_run(pk.add_run(), kicker.upper(), 12, RGBColor(0xE6, 0x9A, 0xC4), bold=True)

    # titulo
    _, tft = textbox(slide, Inches(0.6), Inches(0.52), Inches(11), Inches(0.8))
    pt = tft.paragraphs[0]
    set_run(pt.add_run(), title, 26, WHITE, bold=True)

    # cuerpo
    _, tf = textbox(slide, Inches(0.7), Inches(1.7), Inches(12), Inches(5.2))
    first = True
    for item in bullets:
        text, level, bold = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        prefix = "" if level == 0 else ""
        bullet_char = "\u25B6  " if level == 0 else "\u2013  "
        run = p.add_run()
        color = TUI_BLUE if (level == 0 and bold) else DARK
        if level > 0:
            color = GREY
        set_run(run, bullet_char + text, 18 if level == 0 else 15,
                color, bold=(bold and level == 0))
        p.space_after = Pt(8 if level == 0 else 4)
        if level > 0:
            p.space_after = Pt(4)

    if footer_note:
        band = add_band(slide, Emu(int(SH - Inches(0.75))), Inches(0.75), LIGHT)
        _, tff = textbox(slide, Inches(0.7), Emu(int(SH - Inches(0.72))),
                         Inches(12), Inches(0.65))
        pf = tff.paragraphs[0]
        set_run(pf.add_run(), footer_note, 13, TUI_ACCENT, italic=True)
    return slide


def content_slide_image(number, kicker, title, image_path, caption=None,
                        intro=None):
    """Slide de cabecera + una imagen grande centrada."""
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide, WHITE)
    add_band(slide, 0, Inches(1.35), TUI_BLUE)

    _, tfn = textbox(slide, Inches(11.9), Inches(0.28), Inches(1.2), Inches(0.8))
    pn = tfn.paragraphs[0]
    pn.alignment = PP_ALIGN.RIGHT
    set_run(pn.add_run(), str(number), 30, RGBColor(0x4A, 0x63, 0x9E), bold=True)

    _, tfk = textbox(slide, Inches(0.6), Inches(0.22), Inches(11), Inches(0.4))
    pk = tfk.paragraphs[0]
    set_run(pk.add_run(), kicker.upper(), 12, RGBColor(0xE6, 0x9A, 0xC4), bold=True)

    _, tft = textbox(slide, Inches(0.6), Inches(0.52), Inches(11), Inches(0.8))
    pt = tft.paragraphs[0]
    set_run(pt.add_run(), title, 26, WHITE, bold=True)

    top_img = Inches(1.65)
    if intro:
        _, tfi = textbox(slide, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5))
        pi = tfi.paragraphs[0]
        set_run(pi.add_run(), intro, 15, TUI_BLUE, bold=True)
        top_img = Inches(2.05)

    # imagen centrada, ancho maximo 12"
    from PIL import Image as _PILImage
    with _PILImage.open(str(image_path)) as im:
        w_px, h_px = im.size
    max_w = Inches(12)
    ratio = h_px / w_px
    img_w = max_w
    img_h = Emu(int(max_w * ratio))
    max_h = Inches(4.6)
    if img_h > max_h:
        img_h = max_h
        img_w = Emu(int(max_h / ratio))
    left = Emu(int((SW - img_w) / 2))
    slide.shapes.add_picture(str(image_path), left, top_img,
                             width=img_w, height=img_h)

    if caption:
        band = add_band(slide, Emu(int(SH - Inches(0.7))), Inches(0.7), LIGHT)
        _, tff = textbox(slide, Inches(0.7), Emu(int(SH - Inches(0.68))),
                         Inches(12), Inches(0.6))
        pf = tff.paragraphs[0]
        pf.alignment = PP_ALIGN.CENTER
        set_run(pf.add_run(), caption, 13, TUI_ACCENT, italic=True)
    return slide


def closing_slide(big, small):
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide, TUI_BLUE)
    add_band(slide, Emu(int(SH * 0.5 - Inches(0.06))), Inches(0.08), TUI_ACCENT)
    _, tf = textbox(slide, Inches(1.0), Inches(2.6), Inches(11.3), Inches(2.2))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), big, 32, WHITE, bold=True)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    set_run(p2.add_run(), small, 20, RGBColor(0xC9, 0xD6, 0xF0))
    p2.space_before = Pt(18)
    return slide


# =============================================================================
# SLIDES
# =============================================================================

title_slide(
    "Motor de recomendación turística con IA\ny redistribución de demanda",
    "Personalizacion + sostenibilidad en un mismo motor  ·  Caso TUI",
    "Trabajo Fin de Master · Universidad Complutense de Madrid · Equipo [nombres]",
)

content_slide(
    2, "El problema", "Los recomendadores clasicos agravan el overtourism",
    [
        ("Los sistemas clasicos optimizan solo conversion y satisfaccion", 0, True),
        ("Concentran a todos los viajeros en los destinos ya populares", 1, False),
        ("Consecuencias del overtourism:", 0, True),
        ("Masificacion y presion sobre infraestructuras y recursos", 1, False),
        ("Distribucion desigual de los flujos turisticos y del gasto", 1, False),
        ("Zonas con potencial de crecimiento infrautilizadas", 1, False),
        ("El sistema de recomendacion no es neutral: si solo optimiza clics, agrava el problema", 0, True),
    ],
)

content_slide(
    3, "Objetivo", "Recomendar mejor sin masificar",
    [
        ("Personalizar la experiencia segun intereses y contexto del viajero", 0, True),
        ("Integrar criterios de sostenibilidad en cada recomendacion", 0, True),
        ("Redistribuir de forma inteligente los flujos turisticos", 0, True),
        ("Promover destinos alternativos y zonas menos saturadas", 1, False),
        ("Impulsar temporada baja y productos turisticos emergentes", 1, False),
        ("Doble objetivo: satisfacer al viajero SIN ignorar el impacto en el destino", 0, True),
    ],
)

content_slide_image(
    4, "Metodologia", "Arquitectura: pipeline de extremo a extremo",
    ASSETS / "pipeline.png",
    intro="Cinco bloques encadenados, reproducibles y trazables:",
    caption="No es un notebook aislado: config centralizada, base de datos, API y dashboard real.",
)

content_slide(
    5, "Metodologia + datos", "Datos reales, multi-fuente y multi-idioma",
    [
        ("Scraping: TUI (ES/DE/UK), Booking, TripAdvisor, Google Maps, YouTube", 0, True),
        ("Senal social: Reddit (r/travel, r/solotravel, r/backpacking...)", 0, True),
        ("Estadistica oficial: Eurostat, INE, AEMET (clima) y OpenStreetMap", 0, True),
        ("Cifras clave:", 0, True),
        ("36.063 resenas reales procesadas", 1, False),
        ("39 destinos reales en el catalogo del pipeline", 1, False),
        ("3 mercados TUI: Espana, Alemania y Reino Unido", 1, False),
    ],
    footer_note="Trabajamos con datos reales, no con un dataset de juguete.",
)

content_slide(
    6, "Tecnicas del master (punto clave)", "IA aplicada: NLP + filtrado + scoring propio",
    [
        ("Embeddings semanticos con transformer multilingue (MiniLM, 384 dim)", 0, True),
        ("Fusion semantica paquete + resenas (0.6 / 0.4) y vector hibrido con atributos numericos", 1, False),
        ("Analisis de sentimiento con transformer sobre 36.063 resenas -> satisfaccion por destino", 0, True),
        ("Recomendador: baseline coseno + LightFM (perdida WARP), 64 componentes", 0, True),
        ("Particion train/test 80/20 y semilla fija (reproducibilidad)", 1, False),
    ],
    footer_note="NLP moderno + filtrado colaborativo + scoring propio, todo reproducible.",
)

content_slide(
    7, "Innovacion central", "TDRS · Tourism Demand Redistribution Score",
    [
        ("El corazon diferencial del proyecto: sostenibilidad convertida en numero accionable", 0, True),
        ("Score compuesto por 8 factores ponderados (suma de pesos = 1):", 0, True),
        ("Afinidad · capacidad · accesibilidad · impacto local", 1, False),
        ("Temporada baja · diversificacion · ocupacion · sensibilidad ambiental", 1, False),
        ("Umbral de ocupacion 0,85: a partir de ahi el destino deja de premiarse", 0, True),
        ("Convierte la sostenibilidad en parte del ranking, no en un eslogan", 0, True),
    ],
)

content_slide_image(
    8, "Metodologia", "Re-ranking con escenarios configurables",
    ASSETS / "escenarios.png",
    intro="Mismos candidatos, distinto peso: la tension se parametriza, no se esconde.",
    caption="Redistribucion intensiva: garantiza >=30% de destinos distintos en el top.",
)

content_slide(
    9, "Metodologia + innovacion", "LLM y explicabilidad",
    [
        ("Un LLM (gpt-4o-mini) genera la explicacion en lenguaje natural de cada recomendacion", 0, True),
        ("Control de coste (presupuesto de tokens), validacion de respuesta y fallback a plantillas", 0, True),
        ("Explicabilidad estructurada: desglose del score, motivos, fortalezas y concesiones", 0, True),
        ("El motor no solo recomienda: explica y degrada con elegancia si la IA externa falla", 0, True),
    ],
)

content_slide(
    10, "Resultados (punto clave)", "Un prototipo desplegable, testeado y trazable",
    [
        ("Producto final: dashboard B2B en Streamlit con 4 vistas", 0, True),
        ("Simulador TDRS · Recomendador Espana (API) · Control Web (KPIs) · Datos/Modelo", 1, False),
        ("Demo en vivo: cambiar de escenario y ver como cambia el ranking", 0, True),
        ("Rigor de ingenieria: 141 tests automatizados, pipeline reproducible, API contrastada", 0, True),
        ("Regla de oro: dato ausente != dato estimado. Si falta, se muestra '-', no se inventa", 0, True),
    ],
    footer_note="No es una maqueta: es un prototipo desplegable de punta a punta.",
)

content_slide(
    11, "Impacto y rigor", "Impacto, limites y siguientes pasos",
    [
        ("Impacto de negocio:", 0, True),
        ("Mas satisfaccion del viajero + reparto del gasto hacia zonas con potencial", 1, False),
        ("Impulso de la temporada baja", 1, False),
        ("Limites honestos: coordenadas a sustituir por geocoder; SQLite efimero en cloud; falta pool de conexiones", 0, True),
        ("Siguiente paso: validar con reservas reales de TUI y A/B testing del efecto redistribucion", 0, True),
    ],
)

closing_slide(
    "Recomendar mejor sin masificar",
    "Personalizacion y sostenibilidad en el mismo motor.  Gracias. Preguntas?",
)

out = Path(__file__).resolve().parent.parent / "docs" / "Presentacion_Competicion_Becas.pptx"
out.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(out))
log = out.parent / "_ppt_log.txt"
log.write_text(
    "OK\n%s\nSlides: %d\n" % (out, len(prs.slides._sldIdLst)),
    encoding="utf-8",
)
print("OK ->", out)
print("Slides:", len(prs.slides._sldIdLst))
